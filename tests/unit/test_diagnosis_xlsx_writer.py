from pathlib import Path

import openpyxl

from classifier.models import Diagnosis, DiagnosisExtractionResult
from classifier.output.csv_writer import write_diagnosis_xlsx


def _read_rows(path: Path) -> tuple[list[str], list[tuple]]:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))  # type: ignore[union-attr]
    headers = [str(h) for h in rows[0]]
    body = [tuple("" if v is None else v for v in row) for row in rows[1:]]
    return headers, body


def test_writes_one_row_per_diagnosis(tmp_path: Path) -> None:
    out = tmp_path / "diagnosis_results.xlsx"
    results = [
        DiagnosisExtractionResult(
            file_path=Path("data/1/dr_progress_note.pdf"),
            diagnoses=[
                Diagnosis(name="Type 2 diabetes", icd_code="E11.9"),
                Diagnosis(name="Hypertension", icd_code=None),
            ],
        )
    ]

    write_diagnosis_xlsx(results, out)

    headers, body = _read_rows(out)
    assert headers == ["file_path", "diagnosis", "icd_code"]
    assert body == [
        ("data/1/dr_progress_note.pdf", "Type 2 diabetes", "E11.9"),
        ("data/1/dr_progress_note.pdf", "Hypertension", ""),
    ]


def test_file_with_no_diagnoses_gets_one_empty_row(tmp_path: Path) -> None:
    out = tmp_path / "diagnosis_results.xlsx"
    results = [DiagnosisExtractionResult(file_path=Path("dr_note.pdf"), diagnoses=[])]

    write_diagnosis_xlsx(results, out)

    _, body = _read_rows(out)
    assert body == [("dr_note.pdf", "", "")]


def test_sheet_named_diagnoses(tmp_path: Path) -> None:
    out = tmp_path / "diagnosis_results.xlsx"
    write_diagnosis_xlsx([], out)

    wb = openpyxl.load_workbook(out)
    assert wb.active.title == "Diagnoses"  # type: ignore[union-attr]
