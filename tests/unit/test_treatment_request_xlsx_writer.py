from pathlib import Path

import openpyxl

from classifier.models import TreatmentRequestResult
from classifier.output.csv_writer import write_treatment_request_xlsx


def _read_rows(path: Path) -> tuple[list[str], list[tuple]]:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))  # type: ignore[union-attr]
    headers = [str(h) for h in rows[0]]
    body = [tuple("" if v is None else v for v in row) for row in rows[1:]]
    return headers, body


def test_writes_one_row_per_file(tmp_path: Path) -> None:
    out = tmp_path / "treatment_request_results.xlsx"
    results = [
        TreatmentRequestResult(
            file_path=Path("data/1/dr_progress_note.pdf"),
            treatment_requested=True,
            reasoning="Provider recommended antibiotics",
        ),
        TreatmentRequestResult(
            file_path=Path("data/2/dr_progress_note.pdf"),
            treatment_requested=False,
            reasoning=None,
        ),
    ]

    write_treatment_request_xlsx(results, out)

    headers, body = _read_rows(out)
    assert headers == ["file_path", "treatment_requested", "reasoning"]
    assert body == [
        ("data/1/dr_progress_note.pdf", True, "Provider recommended antibiotics"),
        ("data/2/dr_progress_note.pdf", False, ""),
    ]


def test_sheet_named_treatment_request(tmp_path: Path) -> None:
    out = tmp_path / "treatment_request_results.xlsx"
    write_treatment_request_xlsx([], out)

    wb = openpyxl.load_workbook(out)
    assert wb.active.title == "TreatmentRequest"  # type: ignore[union-attr]
