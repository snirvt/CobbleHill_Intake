import json
from pathlib import Path

import openpyxl

from classifier.models import ErrorCode, ErrorStage, IntakeReviewPipelineResult, RunError
from classifier.output.csv_writer import write_intake_review_xlsx
from classifier.output.examples import intake_review_example
from classifier.output.intake_review_report import (
    intake_review_all_category_groups,
    intake_review_patient_dict,
    intake_review_report,
    intake_review_summary,
)

_ROOT = Path("/root")


def _failed() -> IntakeReviewPipelineResult:
    folder = _ROOT / "FAILED"
    return IntakeReviewPipelineResult(
        patient_folder=folder,
        dr_paths=[folder / "dr_note.pdf"],
        success=False,
        error="age: LLM error: x",
        errors=[
            RunError(stage=ErrorStage.AGE, code=ErrorCode.LLM_FAILED, message="x", folder=folder)
        ],
    )


def _no_dr_note() -> IntakeReviewPipelineResult:
    folder = _ROOT / "ONLY NURSE"
    return IntakeReviewPipelineResult(
        patient_folder=folder, nurse_paths=[folder / "nurse_visit.pdf"], success=True
    )


_REVIEW_FIELDS = [
    "age", "age_years", "patient_group", "verdict", "verdict_phase_1", "verdict_phase_2",
    "explanation", "phase_1", "phase_2",
]


def test_failed_patient_dict_has_identity_error_and_empty_review() -> None:
    out = intake_review_patient_dict(_failed(), _ROOT)
    assert out["patient"] == "FAILED"
    assert out["dr_files"] == ["FAILED/dr_note.pdf"]
    assert out["nurse_files"] == []
    assert out["error"] == "age: LLM error: x"
    assert all(out[k] is None for k in _REVIEW_FIELDS)


def test_no_dr_note_folder_has_identity_and_none_review() -> None:
    out = intake_review_patient_dict(_no_dr_note(), _ROOT)
    assert out["patient"] == "ONLY NURSE"
    assert out["dr_files"] == []
    assert out["nurse_files"] == ["ONLY NURSE/nurse_visit.pdf"]
    assert out["success"] is True and "error" not in out
    assert all(out[k] is None for k in _REVIEW_FIELDS)


def test_summary_counts_verdicts_phases_and_failures() -> None:
    folder_error = RunError(stage=ErrorStage.FOLDER, code=ErrorCode.EMPTY_FOLDER, message="e")
    assert intake_review_summary([_failed(), _no_dr_note()], [folder_error]) == {
        "total": 2, "accept": 0, "reject": 0, "needs_review": 0, "no_nurse_note": 0,
        "no_dr_note": 1, "failed": 1, "phase_1_failed": 0, "phase_2_failed": 0,
        "skipped_folders": 1, "errors": 2,
    }


def test_report_lists_folder_errors_first() -> None:
    folder_error = RunError(
        stage=ErrorStage.FOLDER, code=ErrorCode.EMPTY_FOLDER, message="e", folder=_ROOT / "E"
    )
    report = intake_review_report([_failed()], [folder_error], _ROOT)
    assert list(report) == ["summary", "all_category_groups", "patients", "errors"]
    assert [e["stage"] for e in report["errors"]] == ["folder", "age"]


def test_example_covers_every_outcome() -> None:
    example = intake_review_example()

    assert list(example) == [
        "sharepoint_folder", "web_url", "summary", "all_category_groups", "patients", "errors",
    ]
    patients = example["patients"]
    assert [p["verdict"] for p in patients] == [
        "ACCEPT", "REJECT", "REJECT", "REJECT", "NEEDS_REVIEW", "NO_NURSE_NOTE",
        "NO_NURSE_NOTE", None,
    ]
    assert [p["verdict_phase_1"] for p in patients] == [
        "PASS", "FAIL", "FAIL", "PASS", "PASS", "PASS", "FAIL", None,
    ]
    assert [p["verdict_phase_2"] for p in patients] == [
        "PASS", "NOT_RUN", "NOT_RUN", "FAIL", "UNDECIDED", None, None, None,
    ]
    assert patients[2]["warnings"] == ["age not found — pediatric rules applied"]
    assert patients[0]["phase_2"] == {"dr_nurse_compare": "EXAMPLE - not real data"}
    assert patients[5]["phase_2"] is None and patients[5]["explanation"] is None
    assert example["summary"]["phase_1_failed"] == 3
    text = json.dumps(example)
    assert "EXAMPLE" in text and "/tmp" not in text and '"pairs"' not in text


def test_patient_dict_layout() -> None:
    accepted = intake_review_example()["patients"][0]
    assert list(accepted) == [
        "patient", "dr_files", "nurse_files", "success", "errors", *_REVIEW_FIELDS, "warnings",
    ]
    assert (accepted["age"], accepted["patient_group"]) == ("6 months", "PEDIATRIC")
    assert accepted["phase_1"] == {
        "diagnosis_categories": ["Other", "Respiratory"],
        "diagnoses": [
            {"name": "Example Diagnosis", "icd_code": "Z00.00", "category": "Other"},
            {"name": "Example Respiratory Diagnosis", "icd_code": "Z11.11",
             "category": "Respiratory"},
        ],
        "medical_necessity": True,
        "medical_necessity_reasoning": "EXAMPLE - not real data",
    }


def test_adult_care_fields_are_none() -> None:
    adult = intake_review_example()["patients"][3]
    assert adult["patient_group"] == "ADULT"
    assert adult["phase_1"]["medical_necessity"] is None
    assert adult["phase_1"]["medical_necessity_reasoning"] is None


def test_xlsx_has_a_sheet_per_verdict(tmp_path: Path) -> None:
    from classifier.output import examples

    results = [
        examples._review("A", 40.0, examples.PatientGroup.ADULT, True, None, None),
        _failed(),
        _no_dr_note(),
    ]
    out = tmp_path / "intake.xlsx"

    write_intake_review_xlsx(results, out, local_root=_ROOT)

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == [
        "All", "Accept", "Reject", "Needs Review", "No Nurse Note", "All Categories",
    ]
    header = [c.value for c in wb["All"][1]]
    assert header == [
        "folder", "patient_name", "dr_file_path", "nurse_file_path", "age", "patient_group",
        "verdict", "verdict_phase_1", "verdict_phase_2", "explanation", "diagnosis_categories",
        "diagnoses", "medical_necessity", "medical_necessity_reasoning", "dr_nurse_compare",
        "warnings", "errors",
    ]
    assert wb["All"].max_row == 4
    assert wb["No Nurse Note"].max_row == 2
    assert wb["Accept"].max_row == 1
    rows = list(wb["All"].iter_rows(values_only=True))
    only_nurse = dict(zip(rows[0], rows[3], strict=True))
    assert only_nurse["patient_name"] == "ONLY NURSE"
    assert only_nurse["nurse_file_path"] == "ONLY NURSE/nurse_visit.pdf"
    assert only_nurse["verdict"] is None and only_nurse["dr_file_path"] is None
    assert list(wb["All Categories"].iter_rows(values_only=True)) == [
        ("diagnosis_category", "count"), ("Other", 1), ("Unknown", 1),
    ]


def test_xlsx_row_has_care_and_dr_nurse_compare(tmp_path: Path) -> None:
    from classifier.output import examples

    match = (examples.ClinicalVerdict.MATCH, examples._DIAGNOSIS)
    result = examples._review(
        "P", 0.5, examples.PatientGroup.PEDIATRIC, True, examples.CareCheck.NEEDED, match,
        extra=[examples._SECOND_DIAGNOSIS],
    )
    out = tmp_path / "intake.xlsx"
    write_intake_review_xlsx([result], out)
    rows = list(openpyxl.load_workbook(out)["All"].iter_rows(values_only=True))
    row = dict(zip(rows[0], rows[1], strict=True))
    assert row["verdict"] == "ACCEPT"
    assert row["diagnosis_categories"] == "Other; Respiratory"
    assert row["medical_necessity"] is True
    assert row["medical_necessity_reasoning"] == "EXAMPLE - not real data"
    assert row["dr_nurse_compare"] == "EXAMPLE - not real data"


def test_all_category_groups_use_phase_1_diagnoses() -> None:
    example = intake_review_example()
    assert example["all_category_groups"] == [
        {"diagnosis_category": "Other", "count": 5},
        {"diagnosis_category": "Unknown", "count": 2},
        {"diagnosis_category": "Respiratory", "count": 1},
    ]
    assert intake_review_all_category_groups([_failed(), _no_dr_note()]) == [
        {"diagnosis_category": "Unknown", "count": 1},
    ]
