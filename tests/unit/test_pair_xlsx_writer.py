from pathlib import Path

import openpyxl
import pytest

from classifier.models import (
    CareCheck,
    Diagnosis,
    DiagnosisCheck,
    DocumentMetadata,
    IdentityMatchResult,
    NurseVisitFields,
    PairClassificationResult,
    PairPipelineResult,
    PatientMetadata,
)
from classifier.output.csv_writer import format_diagnoses, resolve_folder, write_pair_xlsx


def _make_pair_result(
    dr: Path,
    nurse: Path,
    overall: str = "MATCH",
    clinical_verdict: str = "MATCH",
    document_metadata: DocumentMetadata | None = None,
    diagnosis_check: DiagnosisCheck | None = None,
    diagnoses: list[Diagnosis] | None = None,
    care_check: CareCheck | None = None,
    care_reasoning: str = "",
) -> PairPipelineResult:
    result = PairClassificationResult(
        dr_file_path=dr,
        nurse_file_path=nurse,
        identity_match=IdentityMatchResult(
            patient_name=True, dob=True, dos=True, sex=True, account_number=True
        ),
        clinical_verdict=clinical_verdict,
        clinical_reasoning="reasoning",
        overall=overall,
        dr_metadata=document_metadata or DocumentMetadata(file_path=dr, raw_text="", meta=PatientMetadata()),
        nurse_fields=NurseVisitFields(),
        diagnosis_check=diagnosis_check,
        diagnoses=diagnoses or [],
        care_check=care_check,
        care_reasoning=care_reasoning,
    )
    return PairPipelineResult(
        dr_file_path=dr, nurse_file_path=nurse, success=True, result=result
    )


def _make_failure(dr: Path, nurse: Path, error: str = "OCR failed") -> PairPipelineResult:
    return PairPipelineResult(
        dr_file_path=dr, nurse_file_path=nurse, success=False, error=error
    )


def _sheet_rows(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[list[object]]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def test_creates_file(tmp_path: Path, sample_document_metadata: DocumentMetadata) -> None:
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([], out)
    assert out.exists()


def test_three_sheets(tmp_path: Path) -> None:
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([], out)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["All", "No Issues", "Issues"]


def test_all_sheet_contains_every_row(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    results = [
        _make_pair_result(dr, nurse, overall="MATCH"),
        _make_pair_result(dr, nurse, overall="MISMATCH"),
        _make_failure(dr, nurse),
    ]
    out = tmp_path / "results.xlsx"
    write_pair_xlsx(results, out)
    wb = openpyxl.load_workbook(out)
    rows = _sheet_rows(wb["All"])
    assert len(rows) == 4  # header + 3 data rows


def test_no_issues_sheet_contains_only_clean_rows(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    clean = _make_pair_result(dr, nurse, overall="MATCH")
    mismatch = _make_pair_result(dr, nurse, overall="MISMATCH")
    failed = _make_failure(dr, nurse)

    out = tmp_path / "results.xlsx"
    write_pair_xlsx([clean, mismatch, failed], out)
    wb = openpyxl.load_workbook(out)
    rows = _sheet_rows(wb["No Issues"])
    assert len(rows) == 2  # header + 1 clean row


def test_issues_sheet_contains_only_issue_rows(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    clean = _make_pair_result(dr, nurse, overall="MATCH")
    mismatch = _make_pair_result(dr, nurse, overall="MISMATCH")
    partial = _make_pair_result(dr, nurse, overall="PARTIAL_MATCH")
    failed = _make_failure(dr, nurse)

    out = tmp_path / "results.xlsx"
    write_pair_xlsx([clean, mismatch, partial, failed], out)
    wb = openpyxl.load_workbook(out)
    rows = _sheet_rows(wb["Issues"])
    assert len(rows) == 4  # header + 3 issue rows


def test_no_issues_and_issues_are_disjoint_and_cover_all(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    results = [
        _make_pair_result(dr, nurse, overall="MATCH"),
        _make_pair_result(dr, nurse, overall="MISMATCH"),
        _make_failure(dr, nurse),
    ]
    out = tmp_path / "results.xlsx"
    write_pair_xlsx(results, out)
    wb = openpyxl.load_workbook(out)
    clean_count = wb["No Issues"].max_row - 1  # subtract header
    issue_count = wb["Issues"].max_row - 1
    assert clean_count + issue_count == len(results)


def test_all_sheets_share_same_headers(tmp_path: Path) -> None:
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([], out)
    wb = openpyxl.load_workbook(out)
    headers = [list(wb[s].iter_rows(max_row=1, values_only=True))[0] for s in wb.sheetnames]
    assert headers[0] == headers[1] == headers[2]


def test_creates_parent_dirs(tmp_path: Path) -> None:
    out = tmp_path / "a" / "b" / "results.xlsx"
    write_pair_xlsx([], out)
    assert out.exists()


def test_failed_row_goes_to_issues(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([_make_failure(dr, nurse, error="timeout")], out)
    wb = openpyxl.load_workbook(out)
    assert wb["Issues"].max_row == 2  # header + 1
    assert wb["No Issues"].max_row == 1  # header only


def test_verbose_false_excludes_detail_columns(tmp_path: Path) -> None:
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([], out, verbose=False)
    wb = openpyxl.load_workbook(out)
    header = list(wb["All"].iter_rows(max_row=1, values_only=True))[0]
    assert "identity_match" not in header
    assert "dr_fields" not in header
    assert "nurse_fields" not in header


def test_verbose_true_includes_detail_columns(tmp_path: Path) -> None:
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([], out, verbose=True)
    wb = openpyxl.load_workbook(out)
    header = list(wb["All"].iter_rows(max_row=1, values_only=True))[0]
    assert "identity_match" in header
    assert "dr_fields" in header
    assert "nurse_fields" in header


_SP_WEB_URL = "https://company.sharepoint.com/sites/Acme/Shared%20Documents/Patient%20Encounters/Notes"


def test_resolve_folder_local(tmp_path: Path) -> None:
    file = tmp_path / "subdir" / "dr_note.pdf"
    assert resolve_folder(file, None, None) == str(tmp_path / "subdir")


def test_resolve_folder_sp_root(tmp_path: Path) -> None:
    file = tmp_path / "dr_note.pdf"
    assert resolve_folder(file, tmp_path, _SP_WEB_URL) == _SP_WEB_URL


def test_resolve_folder_sp_subfolder(tmp_path: Path) -> None:
    file = tmp_path / "PatientX" / "dr_note.pdf"
    result = resolve_folder(file, tmp_path, _SP_WEB_URL)
    assert result == f"{_SP_WEB_URL}/PatientX"


def test_resolve_folder_sp_subfolder_with_spaces(tmp_path: Path) -> None:
    file = tmp_path / "Patient X" / "dr_note.pdf"
    result = resolve_folder(file, tmp_path, _SP_WEB_URL)
    assert result == f"{_SP_WEB_URL}/Patient%20X"


def test_sp_web_url_written_to_xlsx(tmp_path: Path) -> None:
    local_root = tmp_path / "sp_download"
    subfolder = local_root / "PatientX"
    subfolder.mkdir(parents=True)
    dr = subfolder / "dr_note.pdf"
    nurse = subfolder / "nurse_note.pdf"
    result = _make_pair_result(dr, nurse)
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([result], out, local_root=local_root, sp_web_url=_SP_WEB_URL)
    wb = openpyxl.load_workbook(out)
    rows = _sheet_rows(wb["All"])
    header = rows[0]
    data = rows[1]
    folder_val = data[header.index("folder")]
    assert folder_val == f"{_SP_WEB_URL}/PatientX"


def test_patient_name_derived_from_folder_name(tmp_path: Path) -> None:
    patient_folder = tmp_path / "John Doe"
    patient_folder.mkdir()
    dr = patient_folder / "dr_note.pdf"
    nurse = patient_folder / "nurse_note.pdf"
    result = _make_pair_result(dr, nurse)
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([result], out)
    wb = openpyxl.load_workbook(out)
    rows = _sheet_rows(wb["All"])
    header = rows[0]
    data = rows[1]
    assert data[header.index("patient_name")] == "John Doe"


def test_verbose_true_populates_detail_columns(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    result = _make_pair_result(dr, nurse)
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([result], out, verbose=True)
    wb = openpyxl.load_workbook(out)
    rows = _sheet_rows(wb["All"])
    header = rows[0]
    data = rows[1]
    idx = header.index("identity_match")
    assert data[idx] not in ("", None)


# ---------------------------------------------------------------------------
# Diagnosis check columns
# ---------------------------------------------------------------------------

def _row_value(ws: openpyxl.worksheet.worksheet.Worksheet, column: str, row_index: int = 1) -> object:
    rows = _sheet_rows(ws)
    return rows[row_index][list(rows[0]).index(column)]


def test_diagnosis_columns_present_in_both_modes(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    for verbose in (False, True):
        out = tmp_path / f"results_{verbose}.xlsx"
        write_pair_xlsx([_make_pair_result(dr, nurse)], out, verbose=verbose)
        header = list(openpyxl.load_workbook(out)["All"].iter_rows(max_row=1, values_only=True))[0]
        assert "diagnosis_check" in header
        assert "diagnoses" in header


def test_diagnosis_proof_lists_source_and_icd_code(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    result = _make_pair_result(
        dr,
        nurse,
        diagnosis_check=DiagnosisCheck.EXISTS,
        diagnoses=[
            Diagnosis(name="Colic", icd_code="R10.83", source="dr"),
            Diagnosis(name="Reflux", source="nurse"),
        ],
    )
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([result], out)
    ws = openpyxl.load_workbook(out)["All"]
    assert _row_value(ws, "diagnosis_check") == "EXISTS"
    assert _row_value(ws, "diagnoses") == "dr: Colic (R10.83); nurse: Reflux"


def test_missing_diagnosis_row_lands_in_issues_sheet(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    result = _make_pair_result(dr, nurse, overall="MATCH", diagnosis_check=DiagnosisCheck.MISSING)
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([result], out)
    wb = openpyxl.load_workbook(out)
    assert wb["Issues"].max_row == 2  # header + 1
    assert wb["No Issues"].max_row == 1  # header only


def test_skipped_diagnosis_check_is_not_an_issue(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "peds_dr_a.pdf", tmp_path / "peds_nurse_a.pdf"
    result = _make_pair_result(dr, nurse, overall="MATCH", diagnosis_check=None)
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([result], out)
    wb = openpyxl.load_workbook(out)
    assert wb["No Issues"].max_row == 2  # header + 1
    assert not _row_value(wb["All"], "diagnosis_check")  # blank cells read back as None


def test_failed_result_leaves_diagnosis_columns_blank(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "dr_a.pdf", tmp_path / "nurse_a.pdf"
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([_make_failure(dr, nurse)], out)
    ws = openpyxl.load_workbook(out)["All"]
    assert not _row_value(ws, "diagnosis_check")
    assert not _row_value(ws, "diagnoses")


def test_format_diagnoses_empty_list_is_blank() -> None:
    assert format_diagnoses([]) == ""


# ---------------------------------------------------------------------------
# Care check columns
# ---------------------------------------------------------------------------

def test_care_columns_present_in_both_modes(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "peds_dr_a.pdf", tmp_path / "peds_nurse_a.pdf"
    for verbose in (False, True):
        out = tmp_path / f"care_{verbose}.xlsx"
        write_pair_xlsx([_make_pair_result(dr, nurse)], out, verbose=verbose)
        header = list(openpyxl.load_workbook(out)["All"].iter_rows(max_row=1, values_only=True))[0]
        assert "care_check" in header
        assert "care_reasoning" in header


def test_care_needed_row_is_clean(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "peds_dr_a.pdf", tmp_path / "peds_nurse_a.pdf"
    result = _make_pair_result(
        dr, nurse, care_check=CareCheck.NEEDED, care_reasoning="Home nursing ordered."
    )
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([result], out)
    wb = openpyxl.load_workbook(out)
    assert wb["No Issues"].max_row == 2  # header + 1
    assert _row_value(wb["All"], "care_check") == "NEEDED"
    assert _row_value(wb["All"], "care_reasoning") == "Home nursing ordered."


def test_care_not_needed_row_lands_in_issues_sheet(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "peds_dr_a.pdf", tmp_path / "peds_nurse_a.pdf"
    result = _make_pair_result(dr, nurse, overall="MATCH", care_check=CareCheck.NOT_NEEDED)
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([result], out)
    wb = openpyxl.load_workbook(out)
    assert wb["Issues"].max_row == 2  # header + 1
    assert wb["No Issues"].max_row == 1  # header only


def test_skipped_care_check_is_not_an_issue(tmp_path: Path) -> None:
    dr, nurse = tmp_path / "hospital_dr_a.pdf", tmp_path / "hospital_nurse_a.pdf"
    result = _make_pair_result(dr, nurse, overall="MATCH", care_check=None)
    out = tmp_path / "results.xlsx"
    write_pair_xlsx([result], out)
    wb = openpyxl.load_workbook(out)
    assert wb["No Issues"].max_row == 2  # header + 1
    assert not _row_value(wb["All"], "care_check")
