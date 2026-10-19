from pathlib import Path

import openpyxl
import pytest

from classifier.models import (
    CareCheck,
    Diagnosis,
    DiagnosisCategory,
    DiagnosisCheck,
    DrValidityPipelineResult,
    DrValidityResult,
    PatientGroup,
    ValidityVerdict,
)
from classifier.output.csv_writer import (
    format_age,
    patient_display_name,
    write_dr_validity_xlsx,
)

_COLUMNS = [
    "folder", "patient_name", "dr_file_path", "verdict", "reasons", "patient_group", "age",
    "main_diagnosis", "diagnosis_category", "diagnosis_check", "diagnoses", "care_check",
    "care_reasoning", "errors",
]


def _ok(
    root: Path,
    name: str,
    verdict: ValidityVerdict,
    main: str | None = "GERD",
    category: DiagnosisCategory | None = DiagnosisCategory.REFLUX,
    group: PatientGroup = PatientGroup.PEDIATRIC,
) -> DrValidityPipelineResult:
    folder = root / name
    paths = [folder / "dr_a.pdf", folder / "dr_b.pdf"]
    return DrValidityPipelineResult(
        patient_folder=folder,
        dr_paths=paths,
        success=True,
        result=DrValidityResult(
            patient_folder=folder,
            dr_paths=paths,
            age_years=0.5,
            patient_group=group,
            diagnosis_check=DiagnosisCheck.EXISTS,
            diagnoses=[Diagnosis(name="GERD", icd_code="K21.9")],
            care_check=CareCheck.NEEDED,
            care_reasoning="needs home services",
            main_diagnosis=main,
            diagnosis_category=category,
            verdict=verdict,
            reasons=[] if verdict is ValidityVerdict.VALID else ["no medical necessity"],
        ),
    )


def _failed(root: Path, name: str) -> DrValidityPipelineResult:
    folder = root / name
    return DrValidityPipelineResult(
        patient_folder=folder, dr_paths=[folder / "dr_a.pdf"], success=False, error="OCR failed"
    )


def _rows(path: Path, sheet: str) -> list[dict[str, object]]:
    rows = list(openpyxl.load_workbook(path)[sheet].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def test_sheets_and_columns(tmp_path: Path) -> None:
    out = tmp_path / "dr_validity_results.xlsx"
    write_dr_validity_xlsx([], out)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["All", "Valid", "Needs Review", "Groups", "Categories"]
    header = next(wb["All"].iter_rows(max_row=1, values_only=True))
    assert list(header) == _COLUMNS


def test_row_content(tmp_path: Path) -> None:
    out = tmp_path / "dr_validity_results.xlsx"
    results = [_ok(tmp_path, "DOE, JANE", ValidityVerdict.VALID)]
    write_dr_validity_xlsx(results, out, local_root=tmp_path)

    row = _rows(out, "All")[0]
    assert row["patient_name"] == "DOE, JANE"
    assert row["dr_file_path"] == "DOE, JANE/dr_a.pdf; DOE, JANE/dr_b.pdf"
    assert row["verdict"] == "VALID"
    assert row["patient_group"] == "PEDIATRIC"
    assert row["age"] == "6 months"
    assert row["main_diagnosis"] == "GERD"
    assert row["diagnosis_category"] == "Reflux / GERD"
    assert row["diagnoses"] == "GERD (K21.9)"
    assert row["care_check"] == "NEEDED"


def test_valid_and_needs_review_split(tmp_path: Path) -> None:
    out = tmp_path / "dr_validity_results.xlsx"
    results = [
        _ok(tmp_path, "A", ValidityVerdict.VALID),
        _ok(tmp_path, "B", ValidityVerdict.INVALID),
        _ok(tmp_path, "C", ValidityVerdict.UNKNOWN_AGE, group=PatientGroup.UNKNOWN),
        _failed(tmp_path, "D"),
    ]
    write_dr_validity_xlsx(results, out)

    assert [r["patient_name"] for r in _rows(out, "Valid")] == ["A"]
    assert [r["patient_name"] for r in _rows(out, "Needs Review")] == ["B", "C", "D"]
    assert len(_rows(out, "All")) == 4


def test_failed_row_shows_error_and_blank_checks(tmp_path: Path) -> None:
    out = tmp_path / "dr_validity_results.xlsx"
    write_dr_validity_xlsx([_failed(tmp_path, "D")], out)
    row = _rows(out, "All")[0]
    assert row["errors"] == "OCR failed"
    assert not row["verdict"]
    assert not row["main_diagnosis"]


def test_groups_sheet_counts_patients_by_main_diagnosis(tmp_path: Path) -> None:
    out = tmp_path / "dr_validity_results.xlsx"
    results = [
        _ok(tmp_path, "A", ValidityVerdict.VALID, main="Dementia", group=PatientGroup.ADULT),
        _ok(tmp_path, "B", ValidityVerdict.VALID, main="GERD"),
        _ok(tmp_path, "C", ValidityVerdict.INVALID, main="Dementia", group=PatientGroup.ADULT),
        _failed(tmp_path, "D"),
    ]
    write_dr_validity_xlsx(results, out)

    rows = list(openpyxl.load_workbook(out)["Groups"].iter_rows(values_only=True))
    assert rows == [("main_diagnosis", "count"), ("Dementia", 2), ("GERD", 1), ("Unknown", 1)]


# ---------------------------------------------------------------------------
# Patient name when the input is a single patient folder
# ---------------------------------------------------------------------------

def _root_patient(root: Path) -> DrValidityPipelineResult:
    """A patient whose dr notes sit directly in the (temp download) root."""
    return _ok(root.parent, root.name, ValidityVerdict.VALID)


def test_patient_display_name_uses_folder_name(tmp_path: Path) -> None:
    folder = tmp_path / "EISENBERG, SHEINDY"
    assert patient_display_name(folder, tmp_path, "ignored") == "EISENBERG, SHEINDY"


def test_patient_display_name_uses_root_name_for_notes_in_root(tmp_path: Path) -> None:
    root = tmp_path / "cobblehill_sp_x1y2"
    assert patient_display_name(root, root, "EISENBERG, SHEINDY") == "EISENBERG, SHEINDY"


def test_patient_display_name_falls_back_to_root_folder_name(tmp_path: Path) -> None:
    root = tmp_path / "EISENBERG, SHEINDY"
    assert patient_display_name(root, root) == "EISENBERG, SHEINDY"


def test_patient_display_name_resolves_dot_path() -> None:
    assert patient_display_name(Path("."), Path(".")) == Path.cwd().name


def test_single_patient_download_written_with_sharepoint_name(tmp_path: Path) -> None:
    root = tmp_path / "cobblehill_sp_x1y2"
    out = tmp_path / "dr_validity_results.xlsx"

    write_dr_validity_xlsx(
        [_root_patient(root)], out, local_root=root, root_name="EISENBERG, SHEINDY"
    )

    assert _rows(out, "All")[0]["patient_name"] == "EISENBERG, SHEINDY"


# ---------------------------------------------------------------------------
# format_age
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    ("age_years", "expected"),
    [
        (None, ""),
        (-1, ""),
        (0.0, "0 days"),
        (1 / 365.25, "1 day"),
        (0.02, "7 days"),
        (0.04, "2 weeks"),
        (0.1, "5 weeks"),
        (0.17, "2 months"),
        (0.67, "8 months"),
        (1.0, "12 months"),
        (1.99, "23 months"),
        (2.0, "2 years"),
        (17.9, "17 years"),
        (82.4, "82 years"),
    ],
)
def test_format_age_is_human_readable(age_years: float | None, expected: str) -> None:
    assert format_age(age_years) == expected


def test_categories_sheet_counts_patients_by_category(tmp_path: Path) -> None:
    out = tmp_path / "dr_validity_results.xlsx"
    results = [
        _ok(tmp_path, "A", ValidityVerdict.VALID, main="GERD"),
        _ok(tmp_path, "B", ValidityVerdict.VALID, main="Reflux"),
        _ok(
            tmp_path, "C", ValidityVerdict.VALID, main="Dementia",
            category=DiagnosisCategory.DEVELOPMENTAL,
        ),
        _failed(tmp_path, "D"),
    ]
    write_dr_validity_xlsx(results, out)

    rows = list(openpyxl.load_workbook(out)["Categories"].iter_rows(values_only=True))
    assert rows == [
        ("diagnosis_category", "count"),
        ("Reflux / GERD", 2),
        ("Developmental / Cognitive", 1),
        ("Unknown", 1),
    ]
