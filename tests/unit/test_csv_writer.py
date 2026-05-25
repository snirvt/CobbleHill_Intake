from pathlib import Path

import openpyxl
import pytest

from classifier.models import (
    ClassificationResult,
    MedicationData,
    PipelineResult,
)
from classifier.output.csv_writer import write_xlsx


def _make_success(result: ClassificationResult) -> PipelineResult:
    return PipelineResult(file_path=result.file_path, success=True, result=result)


def _make_failure(file_path: Path, error: str) -> PipelineResult:
    return PipelineResult(file_path=file_path, success=False, error=error)


def _read_xlsx(path: Path) -> list[dict[str, object]]:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))  # type: ignore[union-attr]
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, (v if v is not None else "" for v in row))) for row in rows[1:]]


def test_xlsx_creates_file(
    tmp_path: Path,
    sample_document_metadata,  # from conftest
) -> None:
    result = ClassificationResult(
        file_path=sample_document_metadata.file_path,
        category="NEEDS_VISIT",
        task_results={"doctor_visit_needed": "NEEDS_VISIT"},
        metadata=sample_document_metadata,
    )
    out = tmp_path / "out" / "results.xlsx"
    write_xlsx([_make_success(result)], out, task_names=["doctor_visit_needed"])
    assert out.exists()


def test_xlsx_success_row_columns(
    tmp_path: Path,
    sample_document_metadata,
) -> None:
    result = ClassificationResult(
        file_path=sample_document_metadata.file_path,
        category="NEEDS_VISIT",
        task_results={"doctor_visit_needed": "NEEDS_VISIT"},
        metadata=sample_document_metadata,
    )
    out = tmp_path / "results.xlsx"
    write_xlsx([_make_success(result)], out, task_names=["doctor_visit_needed"])

    rows = _read_xlsx(out)
    assert len(rows) == 1
    row = rows[0]
    assert row["success"] is True
    assert row["category"] == "NEEDS_VISIT"
    assert row["doctor_visit_needed"] == "NEEDS_VISIT"
    assert row["patient_name"] == "Test, Patient"
    assert row["dob"] == "05/30/2025"
    assert row["errors"] == ""


def test_xlsx_failure_row(tmp_path: Path, sample_pdf: Path) -> None:
    out = tmp_path / "results.xlsx"
    write_xlsx(
        [_make_failure(sample_pdf, "OCR failed")],
        out,
        task_names=["doctor_visit_needed"],
    )

    rows = _read_xlsx(out)
    assert len(rows) == 1
    row = rows[0]
    assert row["success"] is False
    assert row["category"] == ""
    assert row["doctor_visit_needed"] == ""
    assert row["patient_name"] == ""
    assert row["errors"] == "OCR failed"


def test_xlsx_multiple_rows(
    tmp_path: Path,
    sample_document_metadata,
    sample_pdf: Path,
) -> None:
    success = ClassificationResult(
        file_path=sample_document_metadata.file_path,
        category="NO_VISIT_NEEDED",
        task_results={"doctor_visit_needed": "NO_VISIT_NEEDED"},
        metadata=sample_document_metadata,
    )
    out = tmp_path / "results.xlsx"
    write_xlsx(
        [_make_success(success), _make_failure(sample_pdf, "timeout")],
        out,
        task_names=["doctor_visit_needed"],
    )

    rows = _read_xlsx(out)
    assert len(rows) == 2
    assert rows[0]["category"] == "NO_VISIT_NEEDED"
    assert rows[1]["errors"] == "timeout"


def test_xlsx_creates_parent_dirs(
    tmp_path: Path,
    sample_document_metadata,
) -> None:
    result = ClassificationResult(
        file_path=sample_document_metadata.file_path,
        category="NEEDS_VISIT",
        task_results={},
        metadata=sample_document_metadata,
    )
    deep = tmp_path / "a" / "b" / "c" / "results.xlsx"
    write_xlsx([_make_success(result)], deep, task_names=[])
    assert deep.exists()


def test_xlsx_unknown_task_defaults_empty(
    tmp_path: Path,
    sample_document_metadata,
) -> None:
    result = ClassificationResult(
        file_path=sample_document_metadata.file_path,
        category="NEEDS_VISIT",
        task_results={},  # task_results empty but task_names non-empty
        metadata=sample_document_metadata,
    )
    out = tmp_path / "results.xlsx"
    write_xlsx([_make_success(result)], out, task_names=["doctor_visit_needed"])

    rows = _read_xlsx(out)
    assert rows[0]["doctor_visit_needed"] == ""


def test_xlsx_sheet_named_results(tmp_path: Path) -> None:
    out = tmp_path / "results.xlsx"
    write_xlsx([], out, task_names=[])
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Results"]
