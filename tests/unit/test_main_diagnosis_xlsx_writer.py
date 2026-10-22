from pathlib import Path

import openpyxl

from classifier.models import DiagnosisCategory, MainDiagnosisResult
from classifier.output.csv_writer import write_main_diagnosis_xlsx


def _rows(path: Path, sheet: str) -> list[tuple]:
    ws = openpyxl.load_workbook(path)[sheet]
    return [tuple("" if v is None else v for v in row) for row in ws.iter_rows(values_only=True)]


def _result(
    n: int,
    label: str | None,
    reasoning: str | None = "why",
    category: DiagnosisCategory | None = None,
    icd_code: str | None = None,
) -> MainDiagnosisResult:
    return MainDiagnosisResult(
        file_path=Path(f"data/{n}/dr_progress_note.pdf"),
        main_diagnosis=label,
        diagnosis_category=category,
        icd_code=icd_code,
        reasoning=reasoning,
    )


def test_writes_main_diagnosis_and_groups_sheets(tmp_path: Path) -> None:
    out = tmp_path / "main_diagnosis_results.xlsx"
    write_main_diagnosis_xlsx([], out)
    assert openpyxl.load_workbook(out).sheetnames == ["MainDiagnosis", "Groups", "Categories"]


def test_writes_one_row_per_file(tmp_path: Path) -> None:
    out = tmp_path / "main_diagnosis_results.xlsx"
    results = [
        _result(1, "Pneumonia", "CAP in assessment", DiagnosisCategory.RESPIRATORY, "J18.9"),
        _result(2, None, None),
    ]
    write_main_diagnosis_xlsx(results, out)

    assert _rows(out, "MainDiagnosis") == [
        ("file_path", "main_diagnosis", "main_diagnosis_category", "icd_code", "reasoning"),
        ("data/1/dr_progress_note.pdf", "Pneumonia", "Respiratory", "J18.9", "CAP in assessment"),
        ("data/2/dr_progress_note.pdf", "", "", "", ""),
    ]


def test_groups_sheet_counts_most_common_first(tmp_path: Path) -> None:
    out = tmp_path / "main_diagnosis_results.xlsx"
    results = [
        _result(1, "Asthma"),
        _result(2, "Pneumonia"),
        _result(3, "Pneumonia"),
        _result(4, "UTI"),
        _result(5, "Pneumonia"),
        _result(6, "UTI"),
    ]
    write_main_diagnosis_xlsx(results, out)

    assert _rows(out, "Groups") == [
        ("main_diagnosis", "count"),
        ("Pneumonia", 3),
        ("UTI", 2),
        ("Asthma", 1),
    ]


def test_groups_missing_diagnosis_as_unknown(tmp_path: Path) -> None:
    out = tmp_path / "main_diagnosis_results.xlsx"
    write_main_diagnosis_xlsx([_result(1, None), _result(2, None), _result(3, "Asthma")], out)

    assert ("Unknown", 2) in _rows(out, "Groups")


def test_creates_parent_dirs(tmp_path: Path) -> None:
    out = tmp_path / "nested" / "dir" / "main_diagnosis_results.xlsx"
    write_main_diagnosis_xlsx([_result(1, "Asthma")], out)
    assert out.exists()


def test_categories_sheet_counts_by_category(tmp_path: Path) -> None:
    out = tmp_path / "main_diagnosis_results.xlsx"
    results = [
        _result(1, "GERD", category=DiagnosisCategory.REFLUX),
        _result(2, "Gastric Reflux", category=DiagnosisCategory.REFLUX),
        _result(3, "Hip Fracture", category=DiagnosisCategory.FRACTURE),
        _result(4, None),
    ]
    write_main_diagnosis_xlsx(results, out)

    assert _rows(out, "Categories") == [
        ("main_diagnosis_category", "count"),
        ("Reflux / GERD", 2),
        ("Fracture / Fall", 1),
        ("Unknown", 1),
    ]
