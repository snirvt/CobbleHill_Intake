"""Integration test for the --main-diagnosis CLI flag.

Uses txt notes only (plaintext extractor — no npm/OCR), but hits the configured
ollama model, so it needs that service running. Assertions are model-agnostic:
they check routing and output shape, not the verdict itself.

Run with: pytest -m integration
"""

import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

_REPO_ROOT = Path(__file__).parents[2]


@pytest.mark.integration
def test_cli_main_diagnosis_processes_dr_notes_only(tmp_path: Path) -> None:
    data = tmp_path / "data"
    (data / "1").mkdir(parents=True)
    (data / "1" / "dr_progress_note.txt").write_text("Assessment: acute otitis media.\n")
    (data / "1" / "nurse_visit.txt").write_text("Nurse note text\n")

    result = subprocess.run(
        [sys.executable, "-m", "cli", "--input", str(data), "--main-diagnosis"],
        capture_output=True,
        text=True,
        cwd=_REPO_ROOT,
    )

    assert result.returncode == 0, f"CLI failed:\n{result.stderr}"

    output = json.loads(result.stdout)
    assert len(output) == 1
    assert output[0]["file"].endswith("dr_progress_note.txt")
    assert output[0]["main_diagnosis"] is None or len(output[0]["main_diagnosis"].split()) <= 2

    xlsx_path = _REPO_ROOT / "output" / "main_diagnosis_results.xlsx"
    assert xlsx_path.exists()
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["MainDiagnosis"]
    rows = list(ws.iter_rows(values_only=True))  # type: ignore[union-attr]
    assert rows[0] == ("file_path", "main_diagnosis", "reasoning")
    assert wb.sheetnames == ["MainDiagnosis", "Groups"]
    assert any(str(r[0]).endswith("dr_progress_note.txt") for r in rows[1:])
